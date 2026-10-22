"""Shared spreadsheet formats and column layouts for imports and exports."""

import csv
import math
import re
from io import BytesIO, StringIO
from pathlib import Path

import openpyxl
import xlrd
import xlwt
from django.http import HttpResponse


class SpreadsheetError(ValueError):
    """A file or layout problem that can be shown to the uploader."""


def text_cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(upload):
    extension = Path(getattr(upload, "name", "")).suffix.lower()
    if extension not in {".csv", ".xls", ".xlsx"}:
        raise SpreadsheetError(
            "Unsupported file format. Choose a .csv, .xls, or .xlsx file."
        )
    content = upload.read()
    try:
        if extension == ".csv":
            decoded = content.decode("utf-8-sig")
            if "\x00" in decoded:
                raise SpreadsheetError(
                    "CSV contains binary data; save it as UTF-8 CSV."
                )
            rows = list(csv.reader(StringIO(decoded, newline=""), strict=True))
        elif extension == ".xlsx":
            book = openpyxl.load_workbook(
                BytesIO(content), read_only=True, data_only=True
            )
            try:
                rows = list(book.worksheets[0].values)
            finally:
                book.close()
        else:
            book = xlrd.open_workbook(file_contents=content)
            sheet = book.sheet_by_index(0)
            rows = [sheet.row_values(index) for index in range(sheet.nrows)]
    except UnicodeDecodeError as exc:
        raise SpreadsheetError(
            "Could not decode CSV. Save the file as UTF-8 CSV."
        ) from exc
    except SpreadsheetError:
        raise
    # Readers raise different errors for damaged ZIP, XML and BIFF data.
    except Exception as exc:
        raise SpreadsheetError(
            f"Could not read {extension} file. It may be damaged, empty, "
            "password-protected, or saved with the wrong extension. "
            "Open it in a spreadsheet editor and save it again."
        ) from exc
    rows = [
        (number, [text_cell(cell) for cell in row])
        for number, row in enumerate(rows, 1)
        if any(text_cell(cell) for cell in row)
    ]
    if not rows:
        raise SpreadsheetError(
            "The spreadsheet is empty. Add a header and at least one data row."
        )
    return rows


def normalize_header(value):
    return re.sub(r"[^a-z0-9]", "", value.lower())


# Canonical order is also the exported order. Only required identity/rank columns
# must be present in a named layout; optional omitted columns become blank.
LAYOUTS = {
    "teams": [
        "Team name",
        "School",
        "Hybrid school",
        "Seed",
        "Debater 1 name",
        "Debater 1 status",
        "Debater 1 email",
        "Debater 1 APDA ID",
        "Debater 2 name",
        "Debater 2 status",
        "Debater 2 email",
        "Debater 2 APDA ID",
    ],
    "judges": ["Judge name", "Judge email", "Calibration"],
    "rooms": ["Room name", "Rank"],
    "scratches": ["Team name", "Judge name", "Scratch type"],
}
ROSTER_KINDS = tuple(LAYOUTS)
LAYOUTS.update(
    {"schools": ["School", "APDA ID"], "debaters": ["Debater name", "APDA ID"]}
)
REQUIRED = {
    "schools": {0, 1},
    "debaters": {0, 1},
    "teams": {0, 1, 4, 8},
    "judges": {0, 2},
    "rooms": {0, 1},
    "scratches": {0, 1, 2},
}
ALIASES = {
    "APDA ID": ["id"],
    "Debater name": ["debater", "name"],
    "Team name": ["team"],
    "School": ["school name", "primary school"],
    "Hybrid school": ["hybrid", "hybrid school name"],
    "Seed": ["team seed", "seeding"],
    "Judge name": ["judge", "name"],
    "Judge email": ["email", "email address"],
    "Calibration": ["rank", "judge rank", "judge calibration", "rating"],
    "Room name": ["room", "name"],
    "Rank": ["room rank", "calibration", "rating"],
    "Scratch type": ["type", "scratch"],
}
for _number in (1, 2):
    for _field, _alternates in {
        "name": [f"debater {_number}", f"speaker {_number}", f"speaker {_number} name"],
        "status": [
            f"debater {_number} novice status",
            f"novice {_number}",
            f"debater {_number} novice",
            f"speaker {_number} status",
        ],
        "email": [f"email {_number}", f"speaker {_number} email"],
        "APDA ID": [f"apda id {_number}", f"debater {_number} id"],
    }.items():
        ALIASES[f"Debater {_number} {_field}"] = _alternates


def is_number(value):
    try:
        return math.isfinite(float(value))
    except (ValueError, TypeError, OverflowError):
        return False


def matches_column(kind, index, value):
    label = LAYOUTS[kind][index]
    if not value:
        return index not in REQUIRED[kind]
    if "email" in label.lower():
        return "@" in value
    if label in {"Rank", "Calibration"} or "APDA ID" in label:
        return is_number(value)
    if label == "Seed":
        return value.lower() in {
            "full",
            "full seed",
            "half",
            "half seed",
            "free",
            "free seed",
            "unseeded",
            "un",
            "none",
        }
    if "status" in label:
        return value.lower() in {"n", "nov", "novice", "v", "var", "varsity"}
    if label == "Scratch type":
        return value.lower() in {"team", "tab", "team scratch", "tab scratch"}
    return "@" not in value and not is_number(value)


def judge_extra_columns(headers, data, mapping, round_count):
    """Retain positional school columns and place named rounds chronologically."""
    default_order = all(target == source for target, source in mapping.items())
    remaining = [i for i in range(len(headers)) if i not in mapping.values()]
    rounds = {}
    for index in remaining:
        match = re.fullmatch(
            r"(?:expected)?(?:round|r)(\d+)(?:expected)?",
            normalize_header(headers[index]),
        )
        if match:
            number = int(match.group(1))
            if number < 1 or number > round_count:
                raise SpreadsheetError(
                    f'Column "{headers[index]}" exceeds the configured total of '
                    f"{round_count} rounds."
                )
            if number in rounds:
                raise SpreadsheetError(
                    f"Duplicate expected-round column for round {number}."
                )
            rounds[number] = index
    if not rounds and default_order and round_count and len(remaining) >= round_count:
        candidates = [i for i in remaining if i >= len(LAYOUTS["judges"])][
            -round_count:
        ]
        flags = {
            "",
            "1",
            "true",
            "yes",
            "y",
            "expected",
            "in",
            "x",
            "check",
            "checked",
            "0",
            "false",
            "no",
            "n",
            "out",
        }
        if len(candidates) == round_count and all(
            all(i >= len(row) or row[i].lower() in flags for i in candidates)
            for _, row in data
        ):
            rounds = dict(enumerate(candidates, 1))
    schools = [
        i
        for i in remaining
        if i not in rounds.values()
        and (
            (default_order and i >= len(LAYOUTS["judges"]))
            or normalize_header(headers[i]).startswith("school")
        )
    ]
    return schools + [rounds.get(i) for i in range(1, round_count + 1)]


def column_mapping(kind, rows, round_count=0):
    """Prefer default order, then named columns, then unambiguous value types."""
    labels = LAYOUTS[kind]
    first = rows[0][1]
    aliases = {
        normalize_header(alias): index
        for index, label in enumerate(labels)
        for alias in [label] + ALIASES.get(label, [])
    }
    width = len(labels)
    # Test headerless default data first; a person's literal name may itself be
    # a header alias (for example "Judge" or "Room"). Never drop that first row.
    if all(
        all(
            matches_column(kind, i, row[i] if i < len(row) else "")
            for i in range(width)
        )
        for _, row in rows
    ):
        return (
            rows,
            {i: i for i in range(width)},
            list(range(width, max(len(row) for _, row in rows))),
        )
    mapping = {}
    extras = []
    for index, value in enumerate(first):
        key = aliases.get(normalize_header(value))
        if key is not None:
            if key in mapping:
                raise SpreadsheetError(
                    f'Duplicate column "{labels[key]}" in header row.'
                )
            mapping[key] = index
    typed_header = len(rows) > 1 and any(
        not matches_column(kind, target, first[source])
        and all(
            source < len(row) and matches_column(kind, target, row[source])
            for _, row in rows[1:]
        )
        for target, source in mapping.items()
    )
    has_header = len(mapping) >= 2 or typed_header
    data = rows[1:] if has_header else rows
    if not data:
        raise SpreadsheetError("The spreadsheet has headers but no data rows.")
    if kind == "judges" and has_header:
        extras = judge_extra_columns(first, data, mapping, round_count)
    default_header = all(key == value for key, value in mapping.items())
    default_data = all(
        all(
            matches_column(kind, i, row[i] if i < len(row) else "")
            for i in range(width)
        )
        for _, row in data
    )
    if (not has_header or default_header) and default_data:
        return (
            data,
            {i: i for i in range(width)},
            (
                extras
                if has_header and kind == "judges"
                else list(range(width, len(first)))
            ),
        )
    if not has_header:
        mapping = {}
        extras = []
    # Type inference uses every nonblank data value, never one lucky cell. Text
    # identities cannot be inferred when multiple unlabelled text columns remain.
    used = set(mapping.values()) | set(extras)
    for target in range(width):
        if target in mapping:
            continue
        candidates = [
            source
            for source in range(max(len(row) for _, row in data))
            if source not in used
            and any(source < len(row) and row[source] for _, row in data)
            and all(
                matches_column(kind, target, row[source])
                for _, row in data
                if source < len(row) and row[source]
            )
        ]
        if len(candidates) == 1:
            mapping[target] = candidates[0]
            used.add(candidates[0])
    missing = REQUIRED[kind] - mapping.keys()
    if missing and not has_header:
        # A mostly valid default layout should reach field validation so an
        # invalid value is reported with its column, rather than as a missing header.
        if all(
            len(row) >= width
            and matches_column(kind, 0, row[0])
            and sum(matches_column(kind, i, row[i]) for i in range(width)) >= width - 1
            for _, row in data
        ):
            return data, {i: i for i in range(width)}, list(range(width, len(first)))
    if missing:
        names = ", ".join(f'"{labels[i]}"' for i in sorted(missing))
        raise SpreadsheetError(
            f"Couldn't find an unambiguous column for {names}. "
            "Use the default column order or add the named headers from an export."
        )
    return data, mapping, extras


def spreadsheet_response(kind, extension, headers, rows):
    """Write text cells literally, including strings that look like formulas."""
    if extension == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        content = output.getvalue().encode("utf-8-sig")
        mime = "text/csv; charset=utf-8"
    elif extension == "xlsx":
        book = openpyxl.Workbook()
        sheet = book.active
        for row in [headers] + list(rows):
            sheet.append([text_cell(value) for value in row])
            for cell in sheet[sheet.max_row]:
                cell.data_type = "s"
        output = BytesIO()
        book.save(output)
        content = output.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif extension == "xls":
        book = xlwt.Workbook()
        sheet = book.add_sheet("Data")
        for row_number, row in enumerate([headers] + list(rows)):
            for col, value in enumerate(row):
                sheet.write(row_number, col, text_cell(value))
        output = BytesIO()
        book.save(output)
        content = output.getvalue()
        mime = "application/vnd.ms-excel"
    else:
        raise SpreadsheetError("Unsupported export format. Choose csv, xls, or xlsx.")
    response = HttpResponse(content, content_type=mime)
    response["Content-Disposition"] = f'attachment; filename="{kind}.{extension}"'
    return response
